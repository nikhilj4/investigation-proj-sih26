"""Text Extractor — Extract text from PDF, DOCX, XLSX, CSV, TXT, and images"""

import os
from dataclasses import dataclass, field


@dataclass
class PageContent:
    page_num: int
    text: str


@dataclass
class ExtractedText:
    text: str
    pages: list[PageContent] = field(default_factory=list)
    total_pages: int = 0
    metadata: dict = field(default_factory=dict)


class TextExtractor:
    """Extract text from various document formats."""

    def extract(self, file_path: str, file_type: str) -> ExtractedText:
        """Main extraction dispatcher."""
        file_type = file_type.lower()

        extractors = {
            "pdf": self._extract_pdf,
            "docx": self._extract_docx,
            "doc": self._extract_docx,
            "xlsx": self._extract_xlsx,
            "xls": self._extract_xlsx,
            "csv": self._extract_csv,
            "txt": self._extract_txt,
            "jpg": self._extract_image,
            "jpeg": self._extract_image,
            "png": self._extract_image,
        }

        extractor = extractors.get(file_type)
        if not extractor:
            raise ValueError(f"Unsupported file type: {file_type}")

        return extractor(file_path)

    def _extract_pdf(self, file_path: str) -> ExtractedText:
        """Extract text from PDF using PyMuPDF. Falls back to OCR for scanned pages."""
        import fitz  # PyMuPDF

        doc = fitz.open(file_path)
        pages = []
        full_text_parts = []

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()

            # If very little text found, try OCR
            if len(text) < 50:
                text_ocr = self._ocr_page(page)
                if text_ocr and len(text_ocr) > len(text):
                    text = text_ocr

            if text:
                pages.append(PageContent(page_num=page_num + 1, text=text))
                full_text_parts.append(f"[Page {page_num + 1}]\n{text}")

        doc.close()

        return ExtractedText(
            text="\n\n".join(full_text_parts),
            pages=pages,
            total_pages=len(pages),
        )

    def _ocr_page(self, page) -> str:
        """OCR a single PDF page using pytesseract."""
        try:
            import pytesseract
            from PIL import Image
            import io

            # Render page to image at 300 DPI
            pix = page.get_pixmap(dpi=300)
            img_data = pix.tobytes("png")
            image = Image.open(io.BytesIO(img_data))

            text = pytesseract.image_to_string(image, lang="eng")
            return text.strip()
        except Exception:
            return ""

    def _extract_docx(self, file_path: str) -> ExtractedText:
        """Extract text from DOCX using python-docx."""
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        return ExtractedText(
            text=full_text,
            pages=[PageContent(page_num=1, text=full_text)],
            total_pages=1,
        )

    def _extract_xlsx(self, file_path: str) -> ExtractedText:
        """Extract text from Excel files using openpyxl."""
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages = []
        all_text_parts = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(cell) if cell is not None else "" for cell in row]
                row_text = " | ".join(v for v in row_vals if v)
                if row_text:
                    rows.append(row_text)

            if rows:
                sheet_text = f"[Sheet: {sheet_name}]\n" + "\n".join(rows)
                pages.append(PageContent(page_num=sheet_idx + 1, text=sheet_text))
                all_text_parts.append(sheet_text)

        wb.close()

        return ExtractedText(
            text="\n\n".join(all_text_parts),
            pages=pages,
            total_pages=len(pages),
        )

    def _extract_csv(self, file_path: str) -> ExtractedText:
        """Extract text from CSV files."""
        import pandas as pd

        df = pd.read_csv(file_path, encoding="utf-8", on_bad_lines="skip")
        text_parts = []

        # Header
        text_parts.append(" | ".join(str(c) for c in df.columns))

        # Rows
        for _, row in df.iterrows():
            row_text = " | ".join(str(v) for v in row.values if pd.notna(v))
            if row_text:
                text_parts.append(row_text)

        full_text = "\n".join(text_parts)
        return ExtractedText(
            text=full_text,
            pages=[PageContent(page_num=1, text=full_text)],
            total_pages=1,
            metadata={"rows": len(df), "columns": len(df.columns)},
        )

    def _extract_txt(self, file_path: str) -> ExtractedText:
        """Extract text from plain text files."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

        return ExtractedText(
            text=text,
            pages=[PageContent(page_num=1, text=text)],
            total_pages=1,
        )

    def _extract_image(self, file_path: str) -> ExtractedText:
        """Extract text from images using OCR."""
        import pytesseract
        from PIL import Image

        image = Image.open(file_path)
        text = pytesseract.image_to_string(image, lang="eng").strip()

        return ExtractedText(
            text=text,
            pages=[PageContent(page_num=1, text=text)],
            total_pages=1,
        )


# Singleton
text_extractor = TextExtractor()
